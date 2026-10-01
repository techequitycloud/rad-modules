#!/usr/bin/env python3
# Copyright 2026 Google LLC
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#      http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
"""Regenerates the Cymbal Pools demo assets committed alongside this script.

Maintainer-only: the generated files are committed, so `tofu apply` never runs
this. Re-run it after editing the content below:

    pip install pillow python-docx
    python3 assets/generate_assets.py

Cymbal Pools is a fictional company. All content here is original.
"""

import csv
import datetime
import os
import random

from docx import Document
from docx.shared import Pt
from PIL import Image, ImageDraw, ImageFont

HERE = os.path.dirname(os.path.abspath(__file__))

BROCHURE_PDF = "Cymbal Pools Convention Brochure.pdf"
ANALYSIS_DOCX = "Cymbal Pools Pool Installation Analysis.docx"
POOL_PARTY_PNG = "pool party.png"
PH_TABLE_PNG = "pH table.png"
INSTALL_CSV = "installation_requests.csv"

BROCHURE = [
    ("title", "Cymbal Pools Partner Convention 2026"),
    ("sub", "Riverfront Convention Center, New Orleans - October 14-16"),
    ("h", "Welcome"),
    ("p", "Cymbal Pools brings together installers, service partners and retail "
          "dealers for three days of product launches, hands-on workshops and "
          "certification sessions. This brochure summarises the agenda, the new "
          "product line and the partner programme updates announced this year."),
    ("h", "Day 1 - Product Launches"),
    ("p", "09:00 Keynote: The Connected Backyard. Cymbal Pools unveils AquaSense, "
          "a sensor hub that reports pH, free chlorine, temperature and flow rate "
          "to the Cymbal mobile app every five minutes."),
    ("p", "11:00 TideRunner Robotic Cleaner. Cordless, 4-hour battery, wall and "
          "waterline climbing, and a self-docking lift for easy removal."),
    ("p", "14:00 LumaReef In-Pool Lighting. Colour-tunable LED niches with "
          "music sync and scheduled scenes, rated for saltwater pools."),
    ("h", "Day 2 - Installer Workshops"),
    ("p", "Site inspection and soil assessment; excavation safety; rebar and "
          "shotcrete best practice; plumbing layout for variable-speed pumps; "
          "and a splash pad design clinic for municipal and residential projects."),
    ("h", "Day 3 - Certification"),
    ("p", "Cymbal Certified Installer (CCI) level 1 and level 2 exams are held "
          "in the morning. Partners who certify three or more technicians this "
          "year qualify for Gold tier pricing and priority lead routing."),
    ("h", "Partner Programme Updates"),
    ("p", "Lead routing now uses zip-code coverage maps; warranty claims move to "
          "the partner portal; and co-marketing funds increase to 4 percent of "
          "annual Cymbal product purchases for Gold partners."),
    ("h", "Contact"),
    ("p", "Partner desk: partners@cymbalpools.example - General enquiries: "
          "info@cymbalpools.com"),
]

ANALYSIS = [
    ("title", "Pool Installation Analysis"),
    ("p", "Prepared by Cymbal Pools Operations for the regional installation "
          "teams. This document describes how Cymbal Pools scopes, inspects and "
          "schedules a residential in-ground pool installation."),
    ("h", "1. How we conduct a pool site inspection"),
    ("b", "Confirm property lines, easements and setback requirements with the "
          "local permitting office before visiting the site."),
    ("b", "Call 811 to locate buried utilities; mark gas, electric, water and "
          "irrigation lines within 3 m of the proposed excavation."),
    ("b", "Assess access: the excavator needs a 3 m wide path; note gates, "
          "fences, trees and slopes that must be removed or protected."),
    ("b", "Take soil samples at two depths. Expansive clay or a high water "
          "table requires an engineering review and may add a hydrostatic "
          "relief valve and gravel base."),
    ("b", "Measure the slope across the pool footprint. More than 0.5 m of "
          "fall needs a retaining wall or raised beam design."),
    ("b", "Identify the equipment pad location within 15 m of the pool and "
          "within reach of a dedicated electrical circuit."),
    ("b", "Photograph the site from all four corners and record findings in "
          "the installation request before quoting."),
    ("h", "2. Installation timeline"),
    ("p", "A standard 10 m x 5 m pool takes six to eight weeks: permits (1-2 "
          "weeks), excavation (2 days), steel and plumbing (1 week), shotcrete "
          "and cure (2 weeks), tile and coping (1 week), plaster, fill and "
          "start-up (1 week)."),
    ("h", "3. Common add-ons"),
    ("p", "Hot tubs are requested on roughly a third of installations and add "
          "about two weeks. Waterfalls and rock features add one week. Splash "
          "pads are increasingly requested for families with young children."),
    ("h", "4. Recording a new request"),
    ("p", "Every request is recorded in the installation_requests table with "
          "the pool dimensions in metres, add-on flags, zip code, customer "
          "contact details and the request date."),
]


def _pdf_escape(text):
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _wrap(text, width):
    words, lines, line = text.split(), [], ""
    for w in words:
        if len(line) + len(w) + 1 > width:
            lines.append(line)
            line = w
        else:
            line = f"{line} {w}".strip()
    if line:
        lines.append(line)
    return lines


def write_pdf(path, blocks):
    """Writes a text-based (searchable) PDF using only the standard library.

    A text PDF matters here: Gemini Enterprise's Drive connector indexes the
    text layer, and an image-only PDF would give 'Summarize @file' nothing to
    ground on.
    """
    styles = {"title": (22, "F2", 40), "sub": (12, "F1", 80), "h": (14, "F2", 60), "p": (10.5, "F1", 95)}
    pages, ops, y = [], [], 780
    for kind, text in blocks:
        size, font, width = styles[kind]
        lines = _wrap(text, width)
        if kind == "h":
            y -= 10
        if y - len(lines) * size * 1.4 < 60:
            pages.append(ops)
            ops, y = [], 780
        for line in lines:
            ops.append(f"BT /{font} {size} Tf 56 {y:.1f} Td ({_pdf_escape(line)}) Tj ET")
            y -= size * 1.4
        y -= 6
    pages.append(ops)

    objs = ["<< /Type /Catalog /Pages 2 0 R >>", None,
            "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
            "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold >>"]
    kids = []
    for page_ops in pages:
        stream = "\n".join(page_ops).encode("latin-1")
        objs.append(f"<< /Length {len(stream)} >>\nstream\n{stream.decode('latin-1')}\nendstream")
        content_id = len(objs)
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 842] "
                    f"/Resources << /Font << /F1 3 0 R /F2 4 0 R >> >> /Contents {content_id} 0 R >>")
        kids.append(f"{len(objs)} 0 R")
    objs[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>"

    out, offsets = bytearray(b"%PDF-1.4\n"), []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{body}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    with open(path, "wb") as f:
        f.write(out)


def write_docx(path, blocks):
    doc = Document()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(11)
    for kind, text in blocks:
        if kind == "title":
            doc.add_heading(text, level=0)
        elif kind == "h":
            doc.add_heading(text, level=1)
        elif kind == "b":
            doc.add_paragraph(text, style="List Bullet")
        else:
            doc.add_paragraph(text)
    doc.core_properties.author = "Cymbal Pools Operations"
    doc.save(path)


def _font(size, bold=False):
    candidates = ["/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf",
                  "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"]
    for c in candidates:
        if os.path.exists(c):
            return ImageFont.truetype(c, size)
    return ImageFont.load_default(size)


def write_pool_party(path):
    """Announcement banner for the 'Annual pool party' Gemini Enterprise announcement."""
    w, h = 1200, 630
    img = Image.new("RGB", (w, h), "#0b6fa4")
    d = ImageDraw.Draw(img)
    for i in range(h // 2, h):  # pool water gradient
        t = (i - h // 2) / (h // 2)
        d.line([(0, i), (w, i)], fill=(int(20 + 40 * t), int(160 - 40 * t), int(210 - 30 * t)))
    for k in range(9):  # ripples
        y = h // 2 + 30 + k * 32
        for x in range(0, w, 120):
            d.arc([x + (k % 2) * 60, y, x + 90 + (k % 2) * 60, y + 18], 200, 340, fill="#bfe9ff", width=3)
    d.rectangle([0, h // 2 - 16, w, h // 2], fill="#f4e3c3")  # pool coping
    d.ellipse([930, 40, 1090, 200], fill="#ffd23f")  # sun
    for cx, col in ((180, "#ff5d8f"), (300, "#ffd23f"), (420, "#7bdff2")):  # pool floats
        d.ellipse([cx - 55, 440, cx + 55, 500], outline=col, width=18)
    d.text((70, 70), "Annual Pool Party", font=_font(78, True), fill="white")
    d.text((74, 170), "Join us for a splash!  RSVP today", font=_font(38), fill="#e6f7ff")
    d.text((74, 225), "Cymbal Pools", font=_font(30, True), fill="#ffd23f")
    img.save(path, optimize=True)


def write_ph_table(path):
    """Screenshot-style table for the 'Convert this image to tabular data' demo."""
    rows = [("Reading", "pH", "Water condition", "Action"),
            ("Very low", "6.8", "Acidic - corrosive", "Add sodium carbonate"),
            ("Low", "7.0", "Slightly acidic", "Add sodium bicarbonate"),
            ("Ideal", "7.4", "Balanced", "No action"),
            ("Ideal", "7.6", "Balanced", "No action"),
            ("High", "7.8", "Slightly basic", "Add muriatic acid"),
            ("Very high", "8.2", "Scaling, cloudy water", "Add acid, retest in 6 h")]
    widths = [170, 90, 280, 300]
    rh, pad = 56, 24
    w, h = sum(widths) + 2 * pad, rh * len(rows) + 2 * pad + 60
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    d.text((pad, pad - 4), "Cymbal Pools - Pool Water pH Guide", font=_font(28, True), fill="#0b3954")
    top = pad + 50
    for r, row in enumerate(rows):
        y = top + r * rh
        fill = "#0b6fa4" if r == 0 else ("#eef7fb" if r % 2 else "white")
        d.rectangle([pad, y, w - pad, y + rh], fill=fill, outline="#9cc3d5")
        x = pad
        for c, cell in enumerate(row):
            d.text((x + 12, y + 16), cell, font=_font(22, r == 0), fill="white" if r == 0 else "#1d1d1d")
            x += widths[c]
            d.line([(x, y), (x, y + rh)], fill="#9cc3d5")
    img.save(path, optimize=True)


def write_installation_csv(path, n=60, seed=7):
    rnd = random.Random(seed)
    zips = ["70112", "70115", "70118", "70124", "70005", "70001", "70433", "70458", "70809", "70503"]
    today = datetime.date(2026, 9, 1)
    with open(path, "w", newline="") as f:
        wr = csv.writer(f)
        wr.writerow(["request_id", "request_date", "length_m", "width_m", "depth_m",
                     "has_hot_tub", "has_waterfall", "zip_code", "customer_phone", "customer_email"])
        for i in range(1, n + 1):
            length = rnd.choice([6, 7, 8, 9, 10, 11, 12, 14, 15])
            width = rnd.choice([3, 3.5, 4, 4.5, 5, 6])
            depth = rnd.choice([1.2, 1.5, 1.8, 2.0, 2.4])
            wr.writerow([f"REQ-{1000 + i}", (today - datetime.timedelta(days=rnd.randint(0, 240))).isoformat(),
                         length, width, depth, rnd.random() < 0.33, rnd.random() < 0.2, rnd.choice(zips),
                         f"504.555.{rnd.randint(1000, 9999)}", f"customer{i:03d}@example.com"])


if __name__ == "__main__":
    write_pdf(os.path.join(HERE, BROCHURE_PDF), BROCHURE)
    write_docx(os.path.join(HERE, ANALYSIS_DOCX), ANALYSIS)
    write_pool_party(os.path.join(HERE, POOL_PARTY_PNG))
    write_ph_table(os.path.join(HERE, PH_TABLE_PNG))
    write_installation_csv(os.path.join(HERE, INSTALL_CSV))
    print("assets regenerated in", HERE)
